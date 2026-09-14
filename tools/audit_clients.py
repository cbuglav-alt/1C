#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Аудит справочника клиентов: что не заполнено и где данные противоречивы.

Работает и с нашим справочником, и с выгрузкой из CRM — колонки распознаются
по русским заголовкам, точное соответствие можно задать файлом сопоставления.

    python3 tools/audit_clients.py --file data/clients.sample.csv
    python3 tools/audit_clients.py --file export_crm.csv --map mapping.json
    python3 tools/audit_clients.py --file export_crm.csv --compare data/clients.csv

Выход: отчёт в консоль, подробный список замечаний в out/audit_clients.csv
и сводка в out/audit_clients.md.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys
from collections import Counter, defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Синонимы заголовков: как поле может называться в выгрузке из CRM
SYNONYMS = {
    'client_id': ['client_id', 'id', 'идентификатор', 'код клиента', 'внешний id', 'номер'],
    'name': ['name', 'наименование', 'название', 'клиент', 'организация', 'компания'],
    'inn': ['inn', 'инн'],
    'opf': ['opf', 'опф', 'форма собственности', 'организационно-правовая форма', 'тип клиента'],
    'sno': ['sno', 'сно', 'система налогообложения', 'налоговый режим', 'режим налогообложения'],
    'nds': ['nds', 'ндс', 'плательщик ндс'],
    'employees': ['employees', 'сотрудники', 'количество сотрудников', 'численность', 'штат'],
    'gph': ['gph', 'гпх', 'договоры гпх', 'подрядчики'],
    'kassa': ['kassa', 'касса', 'ккт', 'онлайн-касса'],
    'ved': ['ved', 'вэд', 'внешнеэкономическая деятельность', 'импорт', 'экспорт'],
    'prop': ['prop', 'имущество', 'основные средства', 'транспорт', 'недвижимость'],
    'alco': ['alco', 'алкоголь', 'егаис'],
    'mark': ['mark', 'маркировка', 'прослеживаемость', 'честный знак'],
    'op': ['op', 'обособленные подразделения', 'обособка', 'оп'],
    'accountant': ['accountant', 'бухгалтер', 'ответственный', 'ведущий бухгалтер', 'менеджер'],
    'payroll': ['payroll', 'зарплатный бухгалтер', 'расчетчик', 'расчётчик', 'зарплата'],
    'chief': ['chief', 'главный бухгалтер', 'главбух', 'куратор'],
    'tariff': ['tariff', 'тариф', 'пакет услуг', 'услуга'],
    'advance_day': ['advance_day', 'день аванса', 'аванс'],
    'salary_day': ['salary_day', 'день зарплаты', 'зарплата выплата', 'день выплаты'],
    'status': ['status', 'статус', 'состояние'],
    'start_date': ['start_date', 'дата начала', 'начало обслуживания', 'дата договора'],
}

# Поля, без которых задачи по клиенту не сгенерируются корректно
REQUIRED = ['name', 'opf', 'sno', 'nds', 'employees', 'accountant']

TRUE_WORDS = {'1', 'да', 'true', 'yes', 'есть', 'y', '+', 'применяется'}
FALSE_WORDS = {'0', 'нет', 'false', 'no', 'n', '-', 'отсутствует', 'не применяется'}
KNOWN_SNO = {'ОСНО', 'УСН-Д', 'УСН-ДР', 'ПСН', 'АУСН', 'ЕСХН', 'НПД'}
KNOWN_OPF = {'ООО', 'ИП', 'АО', 'НКО'}


# --------------------------------------------------------------------------- #
def norm(s: str) -> str:
    return re.sub(r'[^a-zа-яё0-9 ]+', ' ', (s or '').strip().lower()).strip()


def build_mapping(headers: list[str], override: dict) -> dict:
    """Сопоставляет заголовки выгрузки с полями нашей модели."""
    mapping, used = {}, set()
    lookup = {norm(h): h for h in headers}
    for field, variants in SYNONYMS.items():
        if field in override:
            mapping[field] = override[field]
            used.add(override[field])
            continue
        for v in variants:
            h = lookup.get(norm(v))
            if h and h not in used:
                mapping[field] = h
                used.add(h)
                break
    return mapping


def read_rows(path: str) -> tuple[list[dict], list[str]]:
    if path.lower().endswith(('.xlsx', '.xlsm')):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit('Для .xlsx нужен openpyxl: pip install openpyxl. '
                             'Либо сохраните выгрузку как CSV.')
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        headers = [str(c) if c is not None else '' for c in next(it)]
        rows = [dict(zip(headers, [('' if c is None else str(c)) for c in r])) for r in it]
        return rows, headers
    with open(path, encoding='utf-8-sig') as f:
        sample = f.read(4096)
        f.seek(0)
        delim = ';' if sample.count(';') >= sample.count(',') else ','
        rdr = csv.DictReader(f, delimiter=delim)
        return [r for r in rdr if any((v or '').strip() for v in r.values())], list(rdr.fieldnames or [])


def get(row: dict, mapping: dict, field: str) -> str:
    col = mapping.get(field)
    return (row.get(col) or '').strip() if col else ''


def tri_bool(v: str):
    """1 / 0 / None — None означает «не заполнено»."""
    s = norm(v)
    if not s:
        return None
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    return None


def inn_valid(inn: str) -> bool:
    if not inn.isdigit():
        return False

    def csum(digits, coeffs):
        return sum(int(d) * c for d, c in zip(digits, coeffs)) % 11 % 10

    if len(inn) == 10:
        return csum(inn, [2, 4, 10, 3, 5, 9, 4, 6, 8]) == int(inn[9])
    if len(inn) == 12:
        return (csum(inn, [7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == int(inn[10])
                and csum(inn, [3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == int(inn[11]))
    return False


# --------------------------------------------------------------------------- #
def audit(rows: list[dict], mapping: dict) -> list[dict]:
    findings = []
    seen_inn = defaultdict(list)

    def add(row_no, name, level, field, message, effect):
        findings.append({'Строка': row_no, 'Клиент': name, 'Уровень': level,
                         'Поле': field, 'Замечание': message, 'Последствие': effect})

    for i, row in enumerate(rows, start=2):
        name = get(row, mapping, 'name') or f'(без наименования, строка {i})'
        status = norm(get(row, mapping, 'status'))
        if status in ('archived', 'архив', 'закрыт', 'расторгнут'):
            continue

        # --- обязательные поля (nds и employees проверяются отдельно, ниже) ---
        for field in REQUIRED:
            if field not in mapping or field in ('nds', 'employees'):
                continue
            if not get(row, mapping, field):
                add(i, name, 'Критично', field,
                    'Поле не заполнено',
                    'Задачи по клиенту сгенерируются неполно или не на того исполнителя')

        opf = get(row, mapping, 'opf').upper()
        sno = get(row, mapping, 'sno').upper().replace(' ', '')
        nds = tri_bool(get(row, mapping, 'nds'))
        emp_raw = get(row, mapping, 'employees')
        emp = int(emp_raw) if emp_raw.isdigit() else None
        inn = re.sub(r'\D', '', get(row, mapping, 'inn'))

        # --- значения из справочника ---
        if opf and opf not in KNOWN_OPF:
            add(i, name, 'Важно', 'opf', f'Незнакомое значение «{opf}»',
                'Правила применимости задач не сработают')
        if sno and sno not in {s.replace(' ', '') for s in KNOWN_SNO}:
            add(i, name, 'Критично', 'sno', f'Незнакомое значение «{sno}»',
                'Налоговый блок задач не сформируется')

        # --- ИНН ---
        if 'inn' in mapping:
            if not inn:
                add(i, name, 'Важно', 'inn', 'ИНН не заполнен',
                    'Невозможна сверка с ФНС и контроль дублей')
            elif not inn_valid(inn):
                add(i, name, 'Критично', 'inn', f'ИНН «{inn}» не проходит проверку контрольной суммы',
                    'Отчётность уйдёт с неверным реквизитом')
            else:
                seen_inn[inn].append((i, name))
                if opf == 'ООО' and len(inn) != 10:
                    add(i, name, 'Важно', 'inn', 'У ООО должно быть 10 цифр ИНН',
                        'Вероятно, перепутаны ОПФ или ИНН')
                if opf == 'ИП' and len(inn) != 12:
                    add(i, name, 'Важно', 'inn', 'У ИП должно быть 12 цифр ИНН',
                        'Вероятно, перепутаны ОПФ или ИНН')

        # --- логические противоречия ---
        if opf == 'ООО' and sno == 'ПСН':
            add(i, name, 'Критично', 'sno', 'Патент применим только для ИП',
                'Неверный налоговый блок задач')
        if sno == 'ОСНО' and nds is False:
            add(i, name, 'Важно', 'nds', 'ОСНО без признака НДС',
                'Не сформируются декларация и уплата НДС — проверьте освобождение по ст. 145 НК РФ')
        if nds is None and 'nds' in mapping:
            add(i, name, 'Критично', 'nds', 'Признак НДС не заполнен',
                'Блок НДС не сформируется вообще')

        if emp is None and 'employees' in mapping:
            add(i, name, 'Критично', 'employees', 'Численность сотрудников не заполнена',
                'Не сформируется весь зарплатный и кадровый блок')
        elif emp and emp > 0:
            if 'payroll' in mapping and not get(row, mapping, 'payroll'):
                add(i, name, 'Важно', 'payroll', 'Есть сотрудники, но не назначен зарплатный бухгалтер',
                    'Зарплатные задачи уйдут ведущему бухгалтеру')
            for f, label in (('advance_day', 'аванса'), ('salary_day', 'зарплаты')):
                if f in mapping and not get(row, mapping, f).strip('0 '):
                    add(i, name, 'Важно', f, f'Не указан день выплаты {label}',
                        'Срок расчёта зарплаты встанет по умолчанию и может нарушить ТК РФ')

        # --- признаки, которые чаще всего забывают ---
        for f, label, effect in (
            ('kassa', 'касса', 'Не будет задачи по кассовой дисциплине и сверке с ОФД'),
            ('ved', 'ВЭД', 'Не будет задач по валютному контролю и курсовым разницам'),
            ('prop', 'имущество', 'Не будет авансов и декларации по имущественным налогам'),
            ('mark', 'маркировка', 'Не будет отчёта по прослеживаемым товарам'),
            ('alco', 'алкоголь', 'Не будет деклараций по обороту алкоголя'),
            ('op', 'обособленные подразделения', 'НДФЛ не будет разнесён по ОКТМО'),
        ):
            if f in mapping and tri_bool(get(row, mapping, f)) is None:
                add(i, name, 'Внимание', f, f'Признак «{label}» не заполнен (пусто ≠ «нет»)',
                    effect)

        if 'chief' in mapping and not get(row, mapping, 'chief'):
            add(i, name, 'Внимание', 'chief', 'Не назначен главный бухгалтер',
                'Не сформируются задачи контроля качества и налогового планирования')

        if 'start_date' in mapping and not get(row, mapping, 'start_date'):
            add(i, name, 'Внимание', 'start_date', 'Не указана дата начала обслуживания',
                'Нельзя отделить наши периоды от периодов предыдущего бухгалтера')

    for inn, items in seen_inn.items():
        if len(items) > 1:
            for i, name in items:
                add(i, name, 'Критично', 'inn',
                    f'Дубль ИНН {inn}: строки ' + ', '.join(str(x[0]) for x in items),
                    'Задачи задвоятся или уйдут разным исполнителям')
    return findings


def compare(rows: list[dict], mapping: dict, ref_path: str) -> list[str]:
    with open(ref_path, encoding='utf-8-sig') as f:
        ref = {(r.get('inn') or '').strip(): r.get('name', '')
               for r in csv.DictReader(f, delimiter=';') if r.get('client_id')}
    cur = {re.sub(r'\D', '', get(r, mapping, 'inn')): get(r, mapping, 'name') for r in rows}
    cur.pop('', None)
    ref.pop('', None)
    out = []
    for inn in sorted(set(cur) - set(ref)):
        out.append(f'В выгрузке есть, в справочнике нет: {cur[inn]} (ИНН {inn})')
    for inn in sorted(set(ref) - set(cur)):
        out.append(f'В справочнике есть, в выгрузке нет: {ref[inn]} (ИНН {inn})')
    return out


# --------------------------------------------------------------------------- #
def main() -> int:
    ap = argparse.ArgumentParser(description='Аудит заполненности справочника клиентов')
    ap.add_argument('--file', required=True, help='CSV или XLSX выгрузка')
    ap.add_argument('--map', help='JSON с явным сопоставлением полей: {"sno": "Режим НО"}')
    ap.add_argument('--compare', help='Сверить состав с нашим справочником (CSV)')
    ap.add_argument('--out', default=os.path.join(ROOT, 'out'))
    args = ap.parse_args()

    rows, headers = read_rows(args.file)
    override = json.load(open(args.map, encoding='utf-8')) if args.map else {}
    mapping = build_mapping(headers, override)

    unmapped = [h for h in headers if h not in mapping.values()]
    missing = [f for f in SYNONYMS if f not in mapping]

    print(f'Строк в выгрузке: {len(rows)}')
    print(f'Распознано полей: {len(mapping)} из {len(SYNONYMS)}')
    if missing:
        print('НЕ НАЙДЕНЫ В ВЫГРУЗКЕ: ' + ', '.join(missing))
    if unmapped:
        print('Колонки без сопоставления: ' + ', '.join(unmapped[:12]) +
              (' …' if len(unmapped) > 12 else ''))
    print()

    findings = audit(rows, mapping)
    levels = Counter(f['Уровень'] for f in findings)
    by_field = Counter(f['Поле'] for f in findings)
    clients_with = len({f['Клиент'] for f in findings})

    print(f'Замечаний: {len(findings)} у {clients_with} клиентов из {len(rows)}')
    for lvl in ('Критично', 'Важно', 'Внимание'):
        if levels[lvl]:
            print(f'  {lvl}: {levels[lvl]}')
    print('\nЧаще всего не заполнено:')
    for field, n in by_field.most_common(10):
        print(f'  {field:14} {n}')

    os.makedirs(args.out, exist_ok=True)
    csv_path = os.path.join(args.out, 'audit_clients.csv')
    with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, ['Строка', 'Клиент', 'Уровень', 'Поле', 'Замечание', 'Последствие'],
                           delimiter=';')
        w.writeheader()
        w.writerows(findings)

    diffs = compare(rows, mapping, args.compare) if args.compare else []

    md = [f'# Аудит справочника клиентов', '',
          f'Файл: `{os.path.basename(args.file)}` · строк: {len(rows)} · '
          f'дата проверки: {dt.date.today().isoformat()}', '',
          f'Замечаний: **{len(findings)}** у **{clients_with}** клиентов из {len(rows)}', '']
    if missing:
        md += ['## Полей нет в выгрузке', '',
               'Эти признаки нигде не хранятся — значит, соответствующие задачи '
               'не могут генерироваться автоматически:', '']
        md += [f'- `{f}`' for f in missing] + ['']
    md += ['## Сводка по полям', '', '| Поле | Замечаний |', '|---|---:|']
    md += [f'| `{f}` | {n} |' for f, n in by_field.most_common()]
    md += ['', '## Замечания по клиентам', '',
           '| Клиент | Уровень | Поле | Замечание | Последствие |', '|---|---|---|---|---|']
    for f in sorted(findings, key=lambda x: (['Критично', 'Важно', 'Внимание'].index(x['Уровень']),
                                             x['Клиент'])):
        md.append(f"| {f['Клиент']} | {f['Уровень']} | `{f['Поле']}` | {f['Замечание']} "
                  f"| {f['Последствие']} |")
    if diffs:
        md += ['', '## Расхождения состава', ''] + [f'- {d}' for d in diffs]

    md_path = os.path.join(args.out, 'audit_clients.md')
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(md))

    print(f'\n{csv_path}\n{md_path}')
    if diffs:
        print(f'\nРасхождения состава: {len(diffs)} — подробности в отчёте')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
