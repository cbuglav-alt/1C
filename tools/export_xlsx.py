#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel с планом месяца: все задачи по клиентам, профиль клиентов, нагрузка сотрудников.

    python3 tools/export_xlsx.py --month 2026-10 --clients data/clients.csv

Берёт out/plan_<месяц>.csv, который пишет generate_plan.py. Нужен openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MONTHS_RU = ['', 'январь', 'февраль', 'март', 'апрель', 'май', 'июнь',
             'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь']
PRIORITY = {'A': 'Высокий', 'B': 'Средний', 'C': 'Низкий'}


def read(path: str) -> list[dict]:
    with open(path, encoding='utf-8-sig') as f:
        return list(csv.DictReader(f, delimiter=';'))


def style(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDE6F0')
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def main() -> int:
    ap = argparse.ArgumentParser(description='План месяца в Excel')
    ap.add_argument('--month', required=True, help='YYYY-MM')
    ap.add_argument('--clients', default=os.path.join(ROOT, 'data/clients.csv'))
    ap.add_argument('--out', default=os.path.join(ROOT, 'out'))
    args = ap.parse_args()

    rows = read(os.path.join(args.out, f'plan_{args.month}.csv'))
    clients = {c['client_id']: c for c in read(args.clients)}
    rows.sort(key=lambda r: (r['client'], r['internal_due'], r['code']))

    wb = Workbook()
    ws = wb.active
    ws.title = 'По клиентам'
    ws.append(['Клиент', 'ИНН', 'Код', 'Задача', 'Категория', 'Ответственный', 'Дата начала',
               'Срок (внутренний)', 'Законный срок', 'Приоритет', 'Период', 'Мин'])
    for r in rows:
        ws.append([r['client'], clients.get(r['client_id'], {}).get('inn', ''), r['code'], r['title'],
                   r['block'], r['assignee'], r['start'], r['internal_due'], r['legal_due'],
                   PRIORITY.get(r['crit'], r['crit']), r['period'], int(r['effort_min'] or 0)])
    style(ws, [36, 14, 12, 66, 16, 30, 11, 11, 11, 10, 16, 6])

    per_client = defaultdict(lambda: [0, 0])
    for r in rows:
        per_client[r['client_id']][0] += 1
        per_client[r['client_id']][1] += int(r['effort_min'] or 0)
    yes = lambda v: 'да' if v == '1' else ''  # noqa: E731
    s = wb.create_sheet('Сводка по клиентам')
    s.append(['Клиент', 'ИНН', 'ОПФ', 'СНО', 'НДС', 'Патент', 'Сотрудники', 'ККТ', 'Нулевщик',
              'Маркетплейс', 'Главный бухгалтер', 'Первичник', 'Бухгалтер ЗП', 'Ответственный',
              'Задач', 'Часов'])
    for cid, c in sorted(clients.items(), key=lambda x: x[1]['name']):
        n, m = per_client[cid]
        s.append([c['name'], c['inn'], c['opf'], c['sno'], yes(c['nds']), yes(c['psn']),
                  yes('1' if c['employees'] not in ('', '0') else '0'), yes(c['kassa']),
                  yes(c.get('zero', '')), yes(c.get('mp', '')), c['accountant'], c['assistant'],
                  c['payroll'], c['manager'], n, round(m / 60, 1)])
    style(s, [36, 14, 6, 9, 6, 7, 11, 6, 9, 11, 28, 28, 28, 28, 7, 7])

    per_person = defaultdict(lambda: [0, 0, set()])
    for r in rows:
        p = per_person[r['assignee']]
        p[0] += 1
        p[1] += int(r['effort_min'] or 0)
        p[2].add(r['client_id'])
    p = wb.create_sheet('По сотрудникам')
    p.append(['Сотрудник', 'Задач', 'Часов', 'Клиентов'])
    for name, (n, m, cs) in sorted(per_person.items(), key=lambda x: -x[1][0]):
        p.append([name, n, round(m / 60, 1), len(cs)])
    style(p, [34, 8, 8, 10])

    y, mth = int(args.month[:4]), int(args.month[5:7])
    path = os.path.join(args.out, f'Задачи_{MONTHS_RU[mth]}_{y}_по_клиентам.xlsx')
    wb.save(path)
    print(f'{len(rows)} задач, {len(clients)} клиентов → {path}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
